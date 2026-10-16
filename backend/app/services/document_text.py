from __future__ import annotations
from ..database import engine
from ..config import KNOWLEDGE_DIR, SUPPORTED_KNOWLEDGE_SUFFIXES, VECTOR_INDEX
from typing import Any
from io import BytesIO
from pathlib import Path

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document as WordDocument
except ImportError:
    WordDocument = None
from sqlalchemy import text

def invalidate_vector_index() -> None:
    """Drop cached text and embeddings after a knowledge document changes."""
    VECTOR_INDEX.unlink(missing_ok=True)

def extract_document_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8-sig", errors="replace")
    if suffix == ".pdf":
        if PdfReader is None:
            raise RuntimeError("PDF support is not installed. Install pypdf in the project environment.")
        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")
        return "\n\n".join(
            f"Page {page_number}: {page.extract_text() or ''}"
            for page_number, page in enumerate(reader.pages, start=1)
        )
    if suffix == ".docx":
        if WordDocument is None:
            raise RuntimeError("Word support is not installed. Install python-docx in the project environment.")
        document = WordDocument(BytesIO(content))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    raise ValueError("Supported files are .pdf, .docx, .txt, and .md.")

def ensure_legacy_document_rows() -> None:
    """Treat files that predate scoped document records as campus-wide sources."""
    KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
    with engine.begin() as connection:
        for path in KNOWLEDGE_DIR.iterdir():
            if path.is_file() and path.suffix.lower() in SUPPORTED_KNOWLEDGE_SUFFIXES:
                connection.execute(text("""
                    INSERT INTO knowledge_documents (filename, scope, subject_id, uploaded_by)
                    VALUES (:filename, 'campus', NULL, NULL)
                    ON CONFLICT (filename) DO NOTHING
                """), {"filename": path.name})

def accessible_document_rows(current_user: dict[str, Any]) -> list[dict[str, Any]]:
    ensure_legacy_document_rows()
    user_id = current_user["id"]
    role = current_user["role"]
    if role == "admin":
        where, params = "TRUE", {}
    elif role == "faculty":
        where = """d.scope IN ('campus', 'faculty') OR (
            d.scope IN ('course', 'submission') AND d.subject_id IN (
                SELECT subject_id FROM faculty_subject_assignments WHERE faculty_user_id = :user_id
            )
        )"""
        params = {"user_id": user_id}
    else:
        where = """d.scope = 'campus' OR (
            d.scope = 'course' AND d.subject_id IN (
                SELECT subject_id FROM student_subject_enrollments WHERE student_user_id = :user_id
            )
        ) OR (
            d.scope IN ('private', 'submission') AND d.uploaded_by = :user_id
        )"""
        params = {"user_id": user_id}

    with engine.connect() as connection:
        rows = connection.execute(text(f"""
            SELECT d.id, d.filename, d.scope, d.subject_id, d.uploaded_by,
                   s.name AS subject_name, s.code AS subject_code,
                   u.full_name AS uploader_name, d.created_at
            FROM knowledge_documents d
            LEFT JOIN subjects s ON s.id = d.subject_id
            LEFT JOIN users u ON u.id = d.uploaded_by
            WHERE ({where})
            ORDER BY d.created_at DESC, d.filename
        """), params).mappings().all()
    return [dict(row) for row in rows if (KNOWLEDGE_DIR / row["filename"]).is_file()]
